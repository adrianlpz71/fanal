"""Importador del Excel de inversiones con un libro sintético (misma estructura, datos
inventados)."""

from datetime import date
from decimal import Decimal as D

import openpyxl
from sqlalchemy import func, select

from app.db import SessionLocal
from app.importers.excel_inversiones import apply, read
from app.models import InvTransaction, PortfolioSnapshot, User
from app.services import inversiones as svc


def workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws["A5"], ws["C5"], ws["A6"], ws["C6"] = "Total aportado", "Valor real cartera", 900, 1260.0
    ws["A10"], ws["B10"], ws["A11"], ws["B11"] = "Objetivo (€)", 3000.0, "Actual (€)", 100.0

    ws = wb.create_sheet("Objetivo")
    ws.append([])
    ws.append(["Categoría", "% objetivo", "% mínimo", "% máximo"])
    ws.append(["Fondos", 0.8, 0.7, 0.9])
    ws.append(["Cripto", 0.2, 0.1, 0.3])

    ws = wb.create_sheet("Fondos")
    ws.append(["Fondo", "% objetivo dentro fondos", "Valor", "%", "Desv", "Estado", "Sug", "Notas"])
    ws.append(["Mundo", 1.0, None, None, None, None, None, "Fondo indexado"])
    ws.append([])
    ws.append(["Aportación destinada a fondos (€)", "Input"])

    ws = wb.create_sheet("Cripto")
    ws.append(["Cripto", "Ticker", "% objetivo dentro cripto"] + [None] * 5 + ["Notas"])
    ws.append(["Moneda", "MON", 1.0])
    ws.append(["Otra", "OTR", 0.0])

    ws = wb.create_sheet("Cartera")
    ws.append(["Activo", "Categoría", "Subcategoría", "Plataforma", "Cantidad", "Precio medio (€)",
               "Precio actual (€)", "Valor aportado (€)", "Valor real (€)", "P/L", "P/L %",
               "Fecha", "Notas"])  # fmt: skip
    ws.append(["Mundo", "Fondos", "Mundo", "Broker", 100.0, 8.0, 10.0, 800, 1000.0])
    ws.append(["Moneda", "Cripto", "Moneda", "Exchange", 0.01, 10000.0, 26000.0, 100, 260.0])
    ws.append([None, None, None, None, None, None, None, 0, 0])

    ws = wb.create_sheet("Historial")
    ws.append(
        ["Mes", "Aportado mes (€)", "Total aportado (€)", "Valor real (€)", "R", "R%", "Notas"]
    )
    ws.append(["2026-04", 900.0, 900, 1000.0])
    ws.append(["2026-05", 0.0, 900, 0.0])
    wb.save(path)


def test_import_is_complete_and_idempotent(make_user, tmp_path):
    make_user("ana@example.com")
    with SessionLocal() as db:
        _check(db, tmp_path)


def _check(db, tmp_path):
    user = db.scalar(select(User).where(User.email == "ana@example.com"))
    path = tmp_path / "inv.xlsx"
    workbook(path)
    moneda_ov = {"coingecko_id": "moneda", "price_provider": "coingecko"}
    ov = {"skip": [], "position_date": "2026-10-01", "assets": {"Moneda": moneda_ov}}
    plan = read(path, ov)
    assert [h.name for h in plan.holdings] == ["Mundo", "Moneda"]
    assert [i.name for i in plan.inner] == ["Mundo", "Moneda", "Otra"]  # se para en la fila vacía

    rep = "\n".join(apply(db, user, plan, ov))
    assert "valor 1260.00 € · coste 900.00 €" in rep
    p = svc.portfolio(db, user.id)
    assert (p.value, p.cost) == (D("1260.00"), D("900.00"))
    by = {c.cls.name: c for c in p.classes}
    assert by["Fondos"].target == D("0.8") and by["Cripto"].status == "ok"
    moneda = next(r for r in p.rows if r.asset.name == "Moneda")
    assert moneda.asset.price_provider == "coingecko" and moneda.asset.type == "cripto"
    otra = next(a for a in svc.assets(db, user.id) if a.name == "Otra")
    assert otra.watchlist
    hist = db.scalars(select(PortfolioSnapshot).where(PortfolioSnapshot.user_id == user.id)).all()
    assert [(s.date, s.value_eur, s.manual) for s in hist] == [
        (date(2026, 4, 30), D("1000.00"), True)
    ]

    # Reimportar no duplica nada
    rep2 = "\n".join(apply(db, user, read(path, ov), ov))
    assert "Objetivos (sin cambios)" in rep2
    n = db.scalar(select(func.count()).select_from(InvTransaction)
                  .where(InvTransaction.user_id == user.id))  # fmt: skip
    assert n == 2
