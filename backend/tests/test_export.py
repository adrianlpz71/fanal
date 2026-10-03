"""Exportación de datos e informe fiscal anual."""

import io
import json
from datetime import date

import openpyxl
import pytest

from tests.conftest import auth_header, enroll
from tests.test_api_gastos import add, setup_gastos

YEAR = date.today().year


@pytest.fixture
def h(client, make_user):
    make_user("ana@example.com")
    hh = auth_header(enroll(client, "ana@example.com"))
    setup_gastos(client, hh, balance="1000.00")
    return hh


def test_export_json_and_xlsx_only_mine_and_no_secrets(client, h, make_user):
    add(client, h, "Cena secreta de Ana", "-30.00")
    make_user("eva@example.com")
    h2 = auth_header(enroll(client, "eva@example.com"))
    setup_gastos(client, h2, balance="5.00")
    add(client, h2, "Cosa de Eva", "-1.00")

    r = client.get("/api/export?format=json", headers=h)
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    d = json.loads(r.content)["data"]
    concepts = {m["concept"] for m in d["movements"]}
    assert "Cena secreta de Ana" in concepts and "Cosa de Eva" not in concepts
    raw = r.text.lower()
    assert "password" not in raw and "totp" not in raw and "refresh" not in raw
    assert d["perfil"][0]["email"] == "ana@example.com"

    x = client.get("/api/export?format=xlsx", headers=h)
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    assert "movements" in wb.sheetnames and "accounts" in wb.sheetnames


def test_tax_report(client, h):
    cls = {c["name"]: c["id"] for c in client.get("/api/inv/classes", headers=h).json()}
    a = client.post(
        "/api/inv/assets",
        headers=h,
        json={
            "name": "Fondo",
            "type": "fondo",
            "isin": "IE0000000001",
            "asset_class_id": cls["Fondos"],
        },
    ).json()
    for d, kind, units, amount in [
        (f"{YEAR - 1}-03-01", "compra", "10", "100"),
        (f"{YEAR}-01-10", "compra", "10", "200"),
        (f"{YEAR}-02-01", "venta", "15", "375"),
    ]:
        r = client.post(
            "/api/inv/transactions",
            headers=h,
            json={
                "asset_id": a["id"],
                "kind": kind,
                "trade_date": d,
                "units": units,
                "amount_eur": amount,
            },
        )
        assert r.status_code == 201, r.text
    add(client, h, "Intereses cuenta remunerada", "2.50", kind="ingreso", date=f"{YEAR}-01-31")
    r = client.get(f"/api/analytics/tax-report?year={YEAR}", headers=h).json()
    (sale,) = r["sales"]
    assert (sale["proceeds"], sale["cost"], sale["gain"]) == ("375.00", "200.00", "175.00")
    assert r["income"][0]["kind"] == "interes_cuenta" and r["income_total"] == "2.50"
    assert r["savings_base"] == "177.50"
    names = {y["name"]: y for y in r["year_end"]}
    assert names["Fondo"]["foreign_hint"] is True  # ISIN irlandés
    assert client.get(f"/api/analytics/tax-report?year={YEAR - 1}", headers=h).json()["sales"] == []
