"""Importador del Excel de gastos con un libro sintético (sin datos reales)."""

from decimal import Decimal as D

import openpyxl
import pytest
from openpyxl.styles import PatternFill
from sqlalchemy import func, select

from app.db import SessionLocal
from app.importers.excel_gastos import GREEN, apply, build_plan, read_workbook
from app.models import InstallmentPlan, Movement, PayCycle, RecurringTemplate, User
from app.services import gastos as svc

G = PatternFill(fill_type="solid", fgColor=GREEN)

# (concepto, importe o fórmula, verde, nota C, celda D)
SHEETS = {
    "Julio 26": (100, 1800, [
        ("Bizum Casa", -200, True, None, None),
        ("Telefono", -10, True, None, None),
        ("Cena", "=-60+20+12", True, "Ana", None),
        ("PayPal Fracción 1", -30, True, None, None),
    ]),
    "Agosto 26": (1630, 1800, [
        ("Bizum Casa", -200, True, None, None),
        ("Telefono", -10, True, None, None),
        ("PayPal Fracción 1", -30, True, None, None),
        ("Zapatillas", -90, True, None, None),
        ("Hamburguesa", -15, False, None, None),  # sin verde en ciclo cerrado → se asume cargado
    ]),
    "Septiembre 26": (3085, 1900, [
        ("Bizum Casa", -200, True, None, None),
        ("Telefono", -10, False, None, None),
        ("PayPal Fracción 1", -30, False, None, None),
        ("PayPal Fracción 2", -90, True, None, None),
        ("Préstamo a Javi", "=10-50", True, "Deuda", "=(250)+B6"),
        ("Nada", 0, False, None, None),
    ]),
    "Octubre 26": (None, 1900, [
        ("Bizum Casa", -200, False, None, None),
        ("Telefono", -10, False, None, None),
        ("PayPal Fracción 1", -30, False, None, None),
        ("PayPal Fracción 2", -90, False, None, None),
        ("Seguro coche", -300, False, None, None),
        ("PayPal Fracción 3", -100, False, "Adelantar Pago", None),
    ]),
}  # fmt: skip

OVERRIDES = {
    "from": "Julio 26",
    "bank": "Banco Test",
    "bank_balance": "4660.00",
    "payday_day": 27,
    "skip_notes": ["Adelantar Pago"],
    "concepts": {"préstamo a javi": "Ocio y social", "zapatillas": "Ropa"},
    "recurring": [
        {"concept": "Bizum Casa", "amount": "-200", "category": "Casa y familia"},
        {"concept": "Telefono", "amount": "-10", "category": "Telefonía e internet"},
    ],
}


@pytest.fixture
def xlsx(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (carry, payroll, rows) in SHEETS.items():
        ws = wb.create_sheet(name)
        ws["A1"], ws["B1"], ws["D1"], ws["E1"] = "Cuenta Start", f"={payroll}+C1", "Total", "=1"
        if carry is not None:
            ws["C1"] = carry
        for i, (concept, amount, green, note, d) in enumerate(rows, start=2):
            ws.cell(i, 1, concept)
            ws.cell(i, 2, amount)
            if green:
                ws.cell(i, 1).fill = G
                ws.cell(i, 2).fill = G
            if note:
                ws.cell(i, 3, note)
            if d:
                ws.cell(i, 4, d)
        if name == "Julio 26":
            ws["L1"] = 5  # columna auxiliar
    p = tmp_path / "gastos.xlsx"
    wb.save(p)
    return p


def _run(xlsx, user_id, commit=True):
    sheets = read_workbook(xlsx, OVERRIDES["skip_notes"])
    plan = build_plan(sheets, OVERRIDES)
    with SessionLocal() as db:
        user = db.get(User, user_id)
        rep = apply(db, user, plan, OVERRIDES)
        db.commit() if commit else db.rollback()
    return "\n".join(rep)


def test_read_workbook(xlsx):
    sheets = read_workbook(xlsx, ["Adelantar Pago"])
    jul, _ago, sep, oct_ = sheets
    assert (jul.payroll, jul.carried, jul.expected_end) == (D(1800), D(100), D(1632))
    cena = next(r for r in jul.rows if r.concept == "Cena")
    assert cena.amount == D(-28) and cena.terms == [D(-60), D(20), D(12)] and cena.note == "Ana"
    assert any("L1" in c for c in jul.aux_cells)
    assert any("Deuda" in c for c in sep.aux_cells)
    assert any("importe 0" in s for s in sep.skipped)
    assert any("Adelantar Pago" in s for s in oct_.skipped)
    assert oct_.carried_blank and not oct_.has_posted


def test_import_creates_cycles_plans_and_matches_excel(xlsx, make_user):
    u = make_user()
    report = _run(xlsx, u.id)
    with SessionLocal() as db:
        cycles = {c.label: c for c in db.scalars(select(PayCycle).where(PayCycle.user_id == u.id))}
        assert {k: c.status for k, c in cycles.items()} == {
            "Julio 26": "closed",
            "Agosto 26": "closed",
            "Septiembre 26": "open",
        }
        cur = cycles["Septiembre 26"]
        s = svc.cycle_summary(db, cur)
        assert s.opening == D(4985)
        assert s.available_now == D(4660)  # = saldo real del banco tras el ajuste de +5
        assert s.expected_end == D(4620)  # E1 del Excel (4615) + ajuste de cuadre (+5)
        assert s.installments == D(120)  # cuota pagada (90) + cuota pendiente del plan (30)
        acc = svc.main_account(db, u.id)
        assert svc.account_balance(db, acc) == D(4660)
        # Descuadre de julio (arrastre real 1630 vs previsto 1632) → ajuste de -2 en julio
        adj = db.scalar(select(Movement).where(Movement.cycle_id == cycles["Julio 26"].id,
                                               Movement.kind == "ajuste"))  # fmt: skip
        assert adj is not None and adj.amount == D(-2)
        # Submovimientos
        cena = db.scalar(
            select(Movement).where(Movement.user_id == u.id, Movement.concept == "Cena")
        )
        assert [ln.amount for ln in cena.lines] == [D(-60), D(20), D(12)]
        # Compras fraccionadas vivas
        plans = {p.provider + str(p.total): p for p in db.scalars(
            select(InstallmentPlan).where(InstallmentPlan.user_id == u.id))}  # fmt: skip
        p30 = plans["paypal120.00"]
        assert [i.status for i in p30.installments] == ["pagada", "pagada", "pendiente",
                                                         "pendiente"]  # fmt: skip
        p90 = plans["paypal270.00"]
        assert p90.description == "Zapatillas"  # 1ª cuota = la compra del ciclo anterior
        assert [i.status for i in p90.installments] == ["pagada", "pagada", "pendiente"]
        # Recurrentes y previstos futuros
        assert db.scalar(select(func.count()).select_from(RecurringTemplate)
                         .where(RecurringTemplate.user_id == u.id)) == 2  # fmt: skip
        seguro = db.scalar(select(Movement).where(Movement.concept == "Seguro coche"))
        assert seguro.status == "planned" and seguro.cycle_id is None
        fc = svc.forecast(db, u.id, 1)[0]
        assert fc.label == "Octubre 26"
        assert (fc.payroll, fc.recurring, fc.installments, fc.other) == (
            D(1900),
            D(-210),
            D(-120),
            D(-300),
        )
    assert "Hamburguesa" in report and "se importa como cargado" in report
    assert "Sin categoría" in report


def test_import_is_idempotent(xlsx, make_user):
    u = make_user()
    _run(xlsx, u.id)
    with SessionLocal() as db:
        before = db.scalar(
            select(func.count()).select_from(Movement).where(Movement.user_id == u.id)
        )
    report = _run(xlsx, u.id)
    with SessionLocal() as db:
        after = db.scalar(
            select(func.count()).select_from(Movement).where(Movement.user_id == u.id)
        )
        plans = db.scalar(select(func.count()).select_from(InstallmentPlan)
                          .where(InstallmentPlan.user_id == u.id))  # fmt: skip
    assert after == before and plans == 2
    assert "0 creados" in report and "ya existía" in report


def test_preview_does_not_write(xlsx, make_user):
    u = make_user()
    _run(xlsx, u.id, commit=False)
    with SessionLocal() as db:
        n = db.scalar(select(func.count()).select_from(Movement).where(Movement.user_id == u.id))
    assert n == 0
