"""Importador del Excel de inversiones (plantilla "Cartera real" de Google Sheets).

Uso (vista previa por defecto; --commit para guardar):
    python -m app.importers.excel_inversiones --email X --file inversiones.xlsx \
        [--overrides overrides.json] [--report informe.md] [--commit]

Qué importa:
  · Cartera  → activos con su posición inicial (participaciones + PMP) fechada HOY: la app mide
               la rentabilidad desde la importación. El precio actual se guarda como precio del día.
  · Objetivo → objetivos macro (objetivo / mínimo / máximo) por categoría.
  · Fondos, Cripto, Acciones → objetivos internos; las filas sin objetivo ni posición pasan a la
               watchlist (seguimiento sin dinero).
  · Historial → puntos manuales del historial (meses con valor real > 0).
  · Dashboard → objetivo del fondo de emergencia (si aún no hay uno en Faro).

Es idempotente: reimportar actualiza la posición inicial, añade una versión de objetivos solo si
cambian y no duplica activos. Lo personal (ISIN, proveedores de precio…) va en un JSON aparte.

Formato del JSON de overrides (todo opcional):
    {"assets": {"<nombre en el Excel>": {"isin": "...", "type": "fondo", "price_provider": "ft",
                                       "price_ref": "...", "coingecko_id": "...", "notes": "..."}},
     "platforms": {"<nombre>": {"kind": "broker", "units_decimals": 4}},
     "skip": ["Otra cripto"], "tolerance_pp": {"Fondos": 1, "Cripto": 2},
     "position_date": "AAAA-MM-DD"}
"""

import argparse
import calendar
import hashlib
import json
import sys
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.money import ZERO, q2
from app.models import Asset, InvTransaction, Platform, PortfolioSnapshot, User
from app.services import gastos as gsvc
from app.services import inversiones as svc

TYPE_BY_CLASS = {"Fondos": "fondo", "Cripto": "cripto", "Acciones": "accion"}


def dec(v: Any) -> Decimal | None:
    if v is None or v == "":
        return None
    if isinstance(v, float):
        # openpyxl da floats: se pasan por repr para no arrastrar ruido binario
        return Decimal(repr(v))
    try:
        return Decimal(str(v).replace(",", "."))
    except ArithmeticError:
        return None


@dataclass
class Holding:
    name: str
    cls: str
    subcategory: str | None
    platform: str | None
    units: Decimal
    avg_cost: Decimal
    price: Decimal | None
    excel_cost: Decimal | None
    excel_value: Decimal | None
    notes: str | None


@dataclass
class Inner:
    name: str
    cls: str
    target: Decimal
    code: str | None = None  # ISIN (acciones) o ticker (cripto)
    sector: str | None = None
    notes: str | None = None


@dataclass
class Plan:
    holdings: list[Holding] = field(default_factory=list)
    macro: dict[str, tuple[Decimal, Decimal, Decimal]] = field(default_factory=dict)
    inner: list[Inner] = field(default_factory=list)
    history: list[tuple[date, Decimal, Decimal | None, str | None]] = field(default_factory=list)
    emergency_target: Decimal | None = None
    emergency_current: Decimal | None = None
    dashboard_value: Decimal | None = None
    dashboard_contributed: Decimal | None = None
    warnings: list[str] = field(default_factory=list)


def _rows(ws, header: str) -> tuple[list[str], list[tuple]]:  # type: ignore[no-untyped-def]
    """Filas bajo la cabecera cuya primera celda es `header`."""
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows):
        if r and isinstance(r[0], str) and r[0].strip() == header:
            head = [str(c).strip() if c is not None else "" for c in r]
            body = []
            for x in rows[i + 1 :]:  # la tabla acaba en la primera fila sin nombre
                if not x or x[0] in (None, ""):
                    break
                body.append(x)
            return head, body
    raise ValueError(f"No encuentro la cabecera '{header}' en la hoja {ws.title}")


def _col(head: list[str], *names: str) -> int:
    for n in names:
        for i, h in enumerate(head):
            if h.lower().startswith(n.lower()):
                return i
    raise ValueError(f"Falta la columna {names[0]}")


def _get(row: tuple, i: int) -> Any:
    return row[i] if i < len(row) else None


def read(path: Path, ov: dict[str, Any]) -> Plan:
    wb = openpyxl.load_workbook(path, data_only=True)
    p = Plan()
    skip = set(ov.get("skip", []))

    head, body = _rows(wb["Cartera"], "Activo")
    c = {k: _col(head, *n) for k, n in {
        "cls": ("Categoría",), "sub": ("Subcategoría",), "plat": ("Plataforma",),
        "units": ("Cantidad",), "pmp": ("Precio medio",), "price": ("Precio actual",),
        "cost": ("Valor aportado",), "value": ("Valor real",), "notes": ("Notas",),
    }.items()}  # fmt: skip
    for r in body:
        name = str(r[0]).strip()
        units, pmp = dec(_get(r, c["units"])), dec(_get(r, c["pmp"]))
        if name in skip or not units or pmp is None:
            continue
        p.holdings.append(Holding(
            name, str(_get(r, c["cls"]) or "").strip(), _get(r, c["sub"]), _get(r, c["plat"]),
            units, pmp, dec(_get(r, c["price"])), dec(_get(r, c["cost"])),
            dec(_get(r, c["value"])), _get(r, c["notes"]),
        ))  # fmt: skip

    head, body = _rows(wb["Objetivo"], "Categoría")
    ti, mi, xi = (
        _col(head, "% objetivo"),
        _col(head, "% mín", "% min"),
        _col(head, "% máx", "% max"),
    )
    for r in body:
        t, lo, hi = dec(_get(r, ti)), dec(_get(r, mi)), dec(_get(r, xi))
        if t is not None:
            p.macro[str(r[0]).strip()] = (
                t,
                lo if lo is not None else t,
                hi if hi is not None else t,
            )

    for sheet, cls, header in (("Fondos", "Fondos", "Fondo"), ("Cripto", "Cripto", "Cripto"),
                               ("Acciones", "Acciones", "Acción")):  # fmt: skip
        if sheet not in wb.sheetnames:
            continue
        head, body = _rows(wb[sheet], header)
        ti = _col(head, "% objetivo")
        ni = _col(head, "Notas")
        code_i = next((i for i, h in enumerate(head) if h in ("Ticker", "Codigo", "Código")), None)
        sec_i = next((i for i, h in enumerate(head) if h == "Sector"), None)
        for r in body:
            name = str(r[0]).strip()
            if name in skip:
                continue
            p.inner.append(Inner(
                name, cls, dec(_get(r, ti)) or ZERO,
                str(_get(r, code_i)).strip() if code_i is not None and _get(r, code_i) else None,
                _get(r, sec_i) if sec_i is not None else None, _get(r, ni),
            ))  # fmt: skip

    if "Historial" in wb.sheetnames:
        head, body = _rows(wb["Historial"], "Mes")
        ci, vi, ni = _col(head, "Total aportado"), _col(head, "Valor real"), _col(head, "Notas")
        for r in body:
            value = dec(_get(r, vi))
            try:
                y, m = (int(x) for x in str(r[0]).split("-"))
            except ValueError:
                continue
            if value and value > 0:
                last = date(y, m, calendar.monthrange(y, m)[1])
                p.history.append((last, value, dec(_get(r, ci)), _get(r, ni)))

    if "Dashboard" in wb.sheetnames:
        ws = wb["Dashboard"]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == "Objetivo (€)":
                    p.emergency_target = dec(ws.cell(cell.row, cell.column + 1).value)
                elif cell.value == "Actual (€)":
                    p.emergency_current = dec(ws.cell(cell.row, cell.column + 1).value)
                elif cell.value == "Valor real cartera":
                    p.dashboard_value = dec(ws.cell(cell.row + 1, cell.column).value)
                elif cell.value == "Total aportado":
                    p.dashboard_contributed = dec(ws.cell(cell.row + 1, cell.column).value)
    return p


def _hash(*parts: str) -> str:
    return hashlib.sha256("|".join(parts).encode()).hexdigest()


def apply(db: Session, user: User, p: Plan, ov: dict[str, Any]) -> list[str]:
    today = date.fromisoformat(ov["position_date"]) if ov.get("position_date") else date.today()
    rep = ["# Importación del Excel de inversiones", ""]
    classes = svc.ensure_classes(db, user.id)
    for name, tol in ov.get("tolerance_pp", {}).items():
        if name in classes:
            classes[name].default_tolerance_pp = Decimal(str(tol))

    # Plataformas
    plats = {x.name: x for x in db.scalars(select(Platform).where(Platform.user_id == user.id))}
    for name in sorted({h.platform for h in p.holdings if h.platform}):
        conf = ov.get("platforms", {}).get(name, {})
        if name not in plats:
            plats[name] = Platform(user_id=user.id, name=name, kind=conf.get("kind", "broker"),
                                   units_decimals=conf.get("units_decimals", 4))  # fmt: skip
            db.add(plats[name])
    db.flush()

    existing = {a.name: a for a in svc.assets(db, user.id, include_archived=True)}

    def upsert(name: str, cls: str, **extra: Any) -> Asset:
        conf = {**extra, **{k: v for k, v in ov.get("assets", {}).get(name, {}).items() if v}}
        a = existing.get(name)
        if a is None:
            a = Asset(
                user_id=user.id, name=name, type=conf.pop("type", TYPE_BY_CLASS.get(cls, "otro"))
            )
            db.add(a)
            existing[name] = a
        else:
            conf.pop("type", None)
        if cls in classes:
            a.asset_class_id = classes[cls].id
        for k, v in conf.items():
            if v is not None and hasattr(a, k):
                setattr(a, k, v)
        return a

    # Posiciones
    rep += ["## Posiciones (fecha de alta: " + today.isoformat() + ")", "",
            "| Activo | Categoría | Plataforma | Participaciones | PMP | Precio | Coste | Valor |",
            "|---|---|---|--:|--:|--:|--:|--:|"]  # fmt: skip
    total_v = total_c = ZERO
    for h in p.holdings:
        a = upsert(h.name, h.cls, notes=h.notes,
                   platform_id=plats[h.platform].id if h.platform else None)  # fmt: skip
        db.flush()
        key = _hash("excel-inv", str(user.id), h.name)
        tx = db.scalar(select(InvTransaction).where(InvTransaction.user_id == user.id,
                                                    InvTransaction.dedupe_hash == key))  # fmt: skip
        if tx is None:
            tx = InvTransaction(user_id=user.id, asset_id=a.id, kind="posicion_inicial",
                                status="liquidada", dedupe_hash=key, external_id="excel")  # fmt: skip
            db.add(tx)
        tx.trade_date, tx.units, tx.avg_cost, tx.amount_eur = today, h.units, h.avg_cost, ZERO
        tx.notes = "Posición inicial importada del Excel"
        if h.price is not None:
            svc.set_price(db, a, today, h.price, "excel")
        cost = h.units * h.avg_cost
        value = h.units * h.price if h.price is not None else cost
        total_c, total_v = total_c + cost, total_v + value
        rep.append(f"| {h.name} | {h.cls} | {h.platform or '—'} | {h.units} | {h.avg_cost} | "
                   f"{h.price if h.price is not None else '—'} | {q2(cost)} | {q2(value)} |")  # fmt: skip
        if h.excel_value is not None and q2(h.excel_value) != q2(value):
            p.warnings.append(
                f"{h.name}: el Excel dice valor {q2(h.excel_value)} y sale {q2(value)}"
            )
    rep += ["", f"**Total**: valor {q2(total_v)} € · coste {q2(total_c)} € · "
                f"P/L {q2(total_v - total_c)} €", ""]  # fmt: skip
    if p.dashboard_value is not None and q2(p.dashboard_value) != q2(total_v):
        p.warnings.append(
            f"Dashboard: valor {q2(p.dashboard_value)} ≠ suma de Cartera {q2(total_v)}"
        )
    if p.dashboard_contributed is not None and q2(p.dashboard_contributed) != q2(total_c):
        p.warnings.append(
            f"Dashboard 'Total aportado' = {q2(p.dashboard_contributed)} €, pero el coste de la "
            f"cartera (participaciones × PMP) es {q2(total_c)} €. Fanal usa el coste por PMP."
        )

    # Objetivos internos y watchlist
    db.flush()
    inner_targets: list[svc.AssetTargetIn] = []
    held = {h.name for h in p.holdings}
    watch = []
    for it in p.inner:
        extra: dict[str, Any] = {"notes": it.notes, "sector": it.sector}
        if it.code and it.cls == "Acciones" and len(it.code) == 12:
            extra["isin"] = it.code
        elif it.code:
            extra["ticker"] = it.code
        if it.target == 0 and it.name not in held:
            a = upsert(it.name, it.cls, watchlist=True, **extra)
            watch.append(it.name)
            continue
        a = upsert(it.name, it.cls, **extra)
        a.watchlist = False
        db.flush()
        inner_targets.append(svc.AssetTargetIn(a.id, it.target))
    db.flush()
    macro_t, inner_t = svc.current_targets(db, user.id)
    macro_in = [
        svc.MacroTargetIn(classes[n].id, t, lo, hi) for n, (t, lo, hi) in p.macro.items()
        if n in classes
    ]  # fmt: skip
    same_macro = len(macro_t) == len(macro_in) and all(
        m.asset_class_id in macro_t
        and (macro_t[m.asset_class_id].target_pct, macro_t[m.asset_class_id].min_pct,
             macro_t[m.asset_class_id].max_pct) == (m.target, m.min, m.max)
        for m in macro_in
    )  # fmt: skip
    same_inner = len(inner_t) == len(inner_targets) and all(
        t.asset_id in inner_t and inner_t[t.asset_id].target_pct == t.target for t in inner_targets
    )
    if not (same_macro and same_inner):
        svc.set_targets(db, user.id, macro_in, inner_targets, today)
        rep.append("## Objetivos (versión nueva desde " + today.isoformat() + ")")
    else:
        rep.append("## Objetivos (sin cambios)")
    rep += ["", "| Categoría | Objetivo | Mín | Máx |", "|---|--:|--:|--:|"]
    for n, (t, lo, hi) in p.macro.items():
        rep.append(f"| {n} | {t * 100:.0f} % | {lo * 100:.0f} % | {hi * 100:.0f} % |")
    rep += ["", "| Activo | Objetivo dentro de su categoría |", "|---|--:|"]
    by_id = {a.id: a.name for a in existing.values()}
    for t in inner_targets:
        rep.append(f"| {by_id.get(t.asset_id, t.asset_id)} | {t.target * 100:.0f} % |")
    rep += ["", f"**Watchlist** ({len(watch)}): {', '.join(watch) or '—'}", ""]

    # Historial manual
    rep += ["## Historial (puntos manuales)", ""]
    for d, value, contributed, notes in p.history:
        if d >= date.today():
            continue
        s = db.scalar(select(PortfolioSnapshot).where(PortfolioSnapshot.user_id == user.id,
                                                      PortfolioSnapshot.date == d))  # fmt: skip
        if s is not None and not s.manual:
            continue
        if s is None:
            s = PortfolioSnapshot(user_id=user.id, date=d, manual=True, net_flow_eur=ZERO)
            db.add(s)
        s.value_eur, s.cost_eur = q2(value), q2(contributed) if contributed else None
        rep.append(f"- {d.isoformat()}: valor {q2(value)} € · aportado "
                   f"{q2(contributed) if contributed else '—'} €{f' ({notes})' if notes else ''}")  # fmt: skip
    rep.append("")

    # Fondo de emergencia
    gs = gsvc.get_settings(db, user.id)
    if p.emergency_target and not gs.get("emergency_target"):
        gsvc.save_settings(db, user.id, {"emergency_target": str(p.emergency_target)})
        rep.append(f"Objetivo del fondo de emergencia: {q2(p.emergency_target)} € (nuevo).")
    elif p.emergency_target:
        rep.append(f"Objetivo del fondo de emergencia: ya había uno ({gs['emergency_target']} €), "
                   "no se cambia.")  # fmt: skip
    e = svc.emergency(db, user.id)
    if e and p.emergency_current is not None and q2(e.current) != q2(p.emergency_current):
        p.warnings.append(f"Refugio: el Excel dice {q2(p.emergency_current)} € y la cuenta "
                          f"refugio de Faro tiene {q2(e.current)} €")  # fmt: skip
    svc.save_settings(db, user.id, {"configured": True})
    db.flush()

    rep += ["", "## Avisos", ""] + ([f"- {w}" for w in p.warnings] or ["- Ninguno"])
    return rep


def main() -> None:
    from app.db import SessionLocal

    ap = argparse.ArgumentParser(prog="app.importers.excel_inversiones")
    ap.add_argument("--email", required=True)
    ap.add_argument("--file", required=True, type=Path)
    ap.add_argument("--overrides", type=Path)
    ap.add_argument("--report", type=Path)
    ap.add_argument("--commit", action="store_true", help="Sin esto, solo vista previa")
    a = ap.parse_args()
    ov = json.loads(a.overrides.read_text(encoding="utf-8")) if a.overrides else {}
    plan = read(a.file, ov)
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.email == a.email.lower()))
        if user is None:
            sys.exit(f"No existe el usuario {a.email}")
        rep = apply(db, user, plan, ov)
        rep.insert(1, "_CONFIRMADO_" if a.commit else "_VISTA PREVIA (no se ha guardado nada)_")
        text = "\n".join(rep) + "\n"
        if a.commit:
            db.commit()
        else:
            db.rollback()
    if a.report:
        a.report.write_text(text, encoding="utf-8")
        print(f"Informe en {a.report}")
    else:
        print(text)


if __name__ == "__main__":
    main()
