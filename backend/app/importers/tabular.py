"""Lectura robusta de ficheros tabulares (CSV/XLSX) exportados por bancos y brokers.

- Codificaciones: UTF-8 (con o sin BOM), UTF-16 (con BOM), Windows-1252 / ISO-8859-1.
- CSV: separador detectado entre ';', ',', tab y '|'.
- XLSX: primera hoja (o la indicada); las celdas se devuelven con su tipo nativo
  (fechas como date/datetime, números como float/int) para no perder información.
- Cabecera: se busca la fila que contiene las columnas esperadas (los bancos suelen poner
  varias filas de títulos antes de la tabla)."""

import csv
import io
import unicodedata
from typing import Any

import openpyxl

Row = list[Any]


class TableError(ValueError):
    pass


def _decode(content: bytes) -> str:
    if content.startswith(b"\xff\xfe") or content.startswith(b"\xfe\xff"):
        return content.decode("utf-16")
    if content.startswith(b"\xef\xbb\xbf"):
        return content[3:].decode("utf-8")
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("cp1252", errors="replace")


def read_table(content: bytes, filename: str, sheet: str | None = None) -> list[Row]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        except Exception as e:
            raise TableError(f"No se pudo leer el Excel: {e}") from None
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith(".xls"):
        raise TableError("Formato .xls antiguo: expórtalo como .xlsx o .csv")
    text = _decode(content)
    sample = text[:4096]
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        delim = ";" if sample.count(";") >= sample.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def norm(s: Any) -> str:
    t = unicodedata.normalize("NFKD", str(s or "").strip().lower())
    return "".join(c for c in t if not unicodedata.combining(c))


def find_header(rows: list[Row], required: list[list[str]], max_scan: int = 40) -> int:
    """Índice de la primera fila que contiene, para cada grupo de `required`, alguna de sus
    palabras (normalizadas). Lanza TableError si no la encuentra."""
    for i, row in enumerate(rows[:max_scan]):
        cells = [norm(c) for c in row]
        if all(any(any(k in c for k in group) for c in cells) for group in required):
            return i
    raise TableError("No encuentro la fila de cabecera (fecha / concepto / importe)")
